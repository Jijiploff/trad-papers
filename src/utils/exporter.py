# exporter.py - Versión con exportación a PDF usando reportlab
#                y LaTeX con escape de caracteres especiales + tablas reales

"""
Módulo de exportación. Genera archivos traducidos en múltiples formatos:
- TXT: texto plano
- DOCX: formato Word con estilo académico
- PDF: formato PDF con estilo académico
- LaTeX: preservando estructura original, con escape de caracteres
  especiales y conversión de tablas HTML (MinerU) a longtable/booktabs
- ZIP: paquete consolidado de múltiples archivos
"""
import os
import re
import zipfile
import tempfile
from html.parser import HTMLParser
from pathlib import Path
from typing import List, Optional, Dict, Any
from datetime import datetime

from src.core.models import Document, FileType
from src.utils.logger import get_logger

logger = get_logger(__name__)


# =============================================================================
# Utilidades para exportación a LaTeX
# =============================================================================

class _LatexHTMLTableParser(HTMLParser):
    """
    Parser HTML minimalista (sin dependencias externas) para extraer filas
    y celdas de las tablas que produce MinerU, del tipo:
    <table><tr><td>...</td><td>...</td></tr></table>
    """

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows: List[List[str]] = []
        self._row: Optional[List[str]] = None
        self._cell_parts: Optional[List[str]] = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th"):
            self._cell_parts = []
        elif tag == "br" and self._cell_parts is not None:
            self._cell_parts.append(" ")

    def handle_endtag(self, tag):
        if tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None
        elif tag in ("td", "th") and self._cell_parts is not None:
            text = re.sub(r"\s+", " ", "".join(self._cell_parts)).strip()
            if self._row is not None:
                self._row.append(text)
            self._cell_parts = None

    def handle_data(self, data):
        if self._cell_parts is not None:
            self._cell_parts.append(data)


# Caracteres que LaTeX interpreta como comandos y deben escaparse.
_LATEX_ESCAPE_RE = re.compile(r'[&%$#_{}~^\\]')
_LATEX_ESCAPE_MAP = {
    '&': r'\&',
    '%': r'\%',
    '$': r'\$',
    '#': r'\#',
    '_': r'\_',
    '{': r'\{',
    '}': r'\}',
    '~': r'\textasciitilde{}',
    '^': r'\textasciicircum{}',
    '\\': r'\textbackslash{}',
}
# Detecta fragmentos matemáticos ($...$) para NO escaparlos como texto normal.
_MATH_SPAN_RE = re.compile(r'\$[^$\n]{1,300}?\$')
# Marca el borde entre el final de una fila markdown ("...| ") y el
# principio de la siguiente ("| ...") cuando los saltos de línea se
# perdieron y todo quedó en un solo párrafo.
_TABLE_ROW_BOUNDARY_RE = re.compile(r'\|\s*\|')


class DocumentExporter:
    """Maneja la exportación de documentos traducidos a varios formatos."""

    def __init__(self, export_dir: str = "exports"):
        self.export_dir = Path(export_dir)
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def _generate_filename(self, document: Document, suffix: str = "", ext: str = "txt") -> str:
        """Genera un nombre de archivo seguro para exportación."""
        base = Path(document.filename).stem
        base = "".join(c if c.isalnum() or c in '-_ ' else '_' for c in base)
        base = base.strip().replace(' ', '_')
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        if suffix:
            return f"{base}_{suffix}_{timestamp}.{ext}"
        return f"{base}_{timestamp}.{ext}"

    def export_to_txt(self, document: Document, output_dir: Optional[Path] = None, original_only: bool = False) -> Path:
        """Exporta la traducción a un archivo de texto plano."""
        output_dir = output_dir or self.export_dir
        filename = self._generate_filename(document, suffix="traducido", ext="txt")
        output_path = output_dir / filename

        lines = []
        lines.append(f"Título original: {document.filename}")
        lines.append(f"Traducido el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append(f"Proveedor: {document.provider_used.value if document.provider_used else 'N/A'}")
        lines.append("=" * 80)
        lines.append("")

        for section in document.sections:
            if section.title:
                lines.append(f"=== {section.title.upper()} ===")
                lines.append("")
            if section.translated_text and not original_only:
                lines.append(section.translated_text)
            else:
                lines.append(section.original_text)
            lines.append("")
            lines.append("")

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write("\n".join(lines))

        logger.info(f"Exportado TXT: {output_path}")
        return output_path

    def export_to_pdf(self, document: Document, output_dir: Optional[Path] = None, original_only: bool = False) -> Path:
        """
        Exporta la traducción a PDF con formato académico.
        Usa reportlab para generar PDF nativo sin dependencias externas.
        """
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether, Table, TableStyle
            from reportlab.lib import colors
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.lib.fonts import addMapping
        except ImportError:
            # Fallback: Intentar convertir DOCX a PDF con docx2pdf
            logger.warning("reportlab no instalado. Intentando convertir DOCX a PDF...")
            return self._export_pdf_via_docx(document, output_dir)

        output_dir = output_dir or self.export_dir
        filename = self._generate_filename(document, suffix="traducido", ext="pdf")
        output_path = output_dir / filename

        # Crear el PDF
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=72,
        )

        styles = getSampleStyleSheet()

        # Crear estilos personalizados
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontSize=16,
            alignment=TA_CENTER,
            spaceAfter=12,
            fontName='Helvetica-Bold',
        )

        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading1'],
            fontSize=14,
            alignment=TA_LEFT,
            spaceAfter=6,
            spaceBefore=12,
            fontName='Helvetica-Bold',
        )

        body_style = ParagraphStyle(
            'CustomBody',
            parent=styles['Normal'],
            fontSize=11,
            alignment=TA_JUSTIFY,
            spaceAfter=6,
            fontName='Helvetica',
        )

        meta_style = ParagraphStyle(
            'MetaStyle',
            parent=styles['Normal'],
            fontSize=9,
            alignment=TA_CENTER,
            spaceAfter=6,
            fontName='Helvetica-Oblique',
            textColor='gray',
        )

        story = []

        # Portada
        story.append(Paragraph(f"<b>{Path(document.filename).stem}</b>", title_style))
        story.append(Spacer(1, 0.25 * inch))
        if original_only:
            document_label = "Documento preprocesado con MinerU"
        else:
            document_label = "Documento Traducido Automáticamente"
        story.append(Paragraph(document_label, meta_style))
        story.append(Spacer(1, 0.1 * inch))
        story.append(Paragraph(f"Fecha: {datetime.now().strftime('%d de %B de %Y')}", meta_style))
        story.append(Spacer(1, 0.1 * inch))
        engine = "MINERU" if original_only else (
            document.provider_used.value.upper() if document.provider_used else "N/A"
        )
        story.append(Paragraph(f"Motor: {engine}", meta_style))
        story.append(PageBreak())

        layout_elements = document.metadata.get("layout_elements", [])
        if layout_elements:
            for element in layout_elements:
                self._append_layout_element(
                    story, element, heading_style, body_style, colors, Table, TableStyle,
                    original_only=original_only,
                )
            doc.build(story)
            logger.info(f"Exportado PDF estructurado: {output_path}")
            return output_path

        # Contenido
        for section in document.sections:
            if not section.translated_text and not section.original_text:
                continue

            text = section.translated_text if section.translated_text else section.original_text

            # Título de sección
            if section.title:
                story.append(Paragraph(f"<b>{section.title}</b>", heading_style))
                story.append(Spacer(1, 0.1 * inch))

            # Párrafos del contenido
            paragraphs = text.split('\n\n')
            for para_text in paragraphs:
                para_text = para_text.strip()
                if not para_text:
                    continue
                # Escapar caracteres especiales para reportlab
                para_text = para_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(para_text, body_style))
                story.append(Spacer(1, 0.05 * inch))

            story.append(Spacer(1, 0.2 * inch))

        # Construir el PDF
        doc.build(story)
        logger.info(f"Exportado PDF: {output_path}")
        return output_path

    @staticmethod
    def _append_layout_element(story, element, heading_style, body_style, colors, Table, TableStyle, original_only=False):
        """Renderiza un elemento sin convertir su estructura en texto corrido."""
        from reportlab.platypus import Paragraph, Spacer

        element_type = element.get("type", "paragraph")
        content = (
            element.get("content", "")
            if original_only
            else element.get("translated_content") or element.get("content", "")
        )
        if not content.strip():
            return

        from xml.sax.saxutils import escape

        if element_type == "section_title":
            story.append(Paragraph(escape(content), heading_style))
            story.append(Spacer(1, 6))
            return

        if element_type == "table":
            rows = []
            for line in content.splitlines():
                if set(line.replace("|", "").strip()) <= {"-", ":", " "}:
                    continue
                cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
                rows.append([Paragraph(escape(cell), body_style) for cell in cells])
            if rows:
                table = Table(rows, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.append(table)
                story.append(Spacer(1, 10))
            return

        if element_type == "equation":
            story.append(Paragraph(f"<font name='Courier'>{escape(content).replace(chr(10), '<br/>')}</font>", body_style))
            story.append(Spacer(1, 8))
            return

        if element_type == "figure":
            story.append(Paragraph(f"<b>[Figura preservada]</b> {escape(content)}", body_style))
            story.append(Spacer(1, 10))
            return

        if element_type == "reference":
            story.append(Paragraph(escape(content), body_style))
            story.append(Spacer(1, 4))
            return

        story.append(Paragraph(escape(content).replace("\n", "<br/>"), body_style))
        story.append(Spacer(1, 6))

    def _export_pdf_via_docx(self, document: Document, output_dir: Optional[Path] = None) -> Path:
        """
        Método alternativo: Exporta a DOCX y luego convierte a PDF usando docx2pdf.
        Requiere: pip install docx2pdf (necesita Microsoft Word o LibreOffice)
        """
        try:
            # Primero exportar a DOCX
            docx_path = self.export_to_docx(document, output_dir)

            # Convertir a PDF usando docx2pdf
            from docx2pdf import convert
            pdf_path = docx_path.with_suffix('.pdf')
            convert(str(docx_path), str(pdf_path))

            logger.info(f"Exportado PDF (vía DOCX): {pdf_path}")
            return pdf_path

        except ImportError:
            raise ImportError(
                "Para exportar a PDF necesitas instalar:\n"
                "  - Opción 1 (ligera): pip install reportlab\n"
                "  - Opción 2 (con Word): pip install docx2pdf"
            )
        except Exception as e:
            logger.error(f"Error convirtiendo a PDF: {e}")
            # Fallback: guardar como DOCX y advertir
            logger.warning("No se pudo generar PDF. Se exportó como DOCX en su lugar.")
            return docx_path

    def export_to_docx(self, document: Document, output_dir: Optional[Path] = None, original_only: bool = False) -> Path:
        """
        Exporta la traducción a DOCX con formato académico estándar.
        Incluye: portada, encabezados, cuerpo con estilos, citas preservadas.
        """
        try:
            from docx import Document as DocxDocument
            from docx.shared import Pt, Inches, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
        except ImportError:
            raise ImportError("python-docx es requerido para exportar a DOCX. Instálalo con: pip install python-docx")

        output_dir = output_dir or self.export_dir
        filename = self._generate_filename(document, suffix="traducido", ext="docx")
        output_path = output_dir / filename

        doc = DocxDocument()

        # Configurar estilos base
        style = doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)

        # Márgenes
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.25)

        # Título del documento
        title = doc.add_heading(Path(document.filename).stem, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.bold = True

        # Metadatos
        meta_para = doc.add_paragraph()
        meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta_para.add_run(
            f"\nDocumento traducido automáticamente\n"
            f"Fecha: {datetime.now().strftime('%d de %B de %Y')}\n"
            f"Motor de traducción: {document.provider_used.value.upper() if document.provider_used else 'N/A'}"
        )
        meta_run.font.name = 'Times New Roman'
        meta_run.font.size = Pt(10)
        meta_run.italic = True
        meta_run.font.color.rgb = RGBColor(100, 100, 100)

        doc.add_page_break()

        # Contenido por secciones
        for section in document.sections:
            if not section.translated_text and not section.original_text:
                continue

            text = (
                section.original_text
                if original_only
                else section.translated_text if section.translated_text else section.original_text
            )

            # Determinar nivel de encabezado
            section_type = section.section_type.value
            heading_map = {
                'title': 0,
                'abstract': 1,
                'introduction': 1,
                'methodology': 1,
                'results': 1,
                'discussion': 1,
                'conclusions': 1,
                'references': 1,
                'acknowledgments': 1,
                'appendix': 1,
                'other': 2,
            }
            level = heading_map.get(section_type, 2)

            if section.title:
                heading = doc.add_heading(section.title, level=level)
                for run in heading.runs:
                    run.font.name = 'Times New Roman'

            # Párrafos del contenido
            paragraphs = text.split('\n\n')
            for para_text in paragraphs:
                para_text = para_text.strip()
                if not para_text:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Inches(0.5)
                run = p.add_run(para_text)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)

        doc.save(str(output_path))
        logger.info(f"Exportado DOCX: {output_path}")
        return output_path

    # =========================================================================
    # LaTeX: helpers de escape y de conversión de tablas HTML
    # =========================================================================

    @staticmethod
    def _clean_math_spacing(math_text: str) -> str:
        """
        MinerU suele emitir fórmulas con espacios sobrantes, p. ej.
        '$P _ { 1 }$' en vez de '$P_{1}$'. Esto las normaliza sin tocar
        el contenido matemático en sí.
        """
        inner = math_text[1:-1]
        inner = re.sub(r'\s*_\s*\{\s*', '_{', inner)
        inner = re.sub(r'\s*\^\s*\{\s*', '^{', inner)
        inner = re.sub(r'\s*\}\s*', '}', inner)
        inner = re.sub(r'\s+', ' ', inner).strip()
        return f"${inner}$"

    @classmethod
    def _escape_latex(cls, text: str) -> str:
        """
        Escapa caracteres especiales de LaTeX (& % $ # _ { } ~ ^ \\) en texto
        plano, preservando (y limpiando el espaciado de) los fragmentos
        matemáticos $...$ tal cual, sin escaparlos.
        """
        if not text:
            return text

        def _escape_plain(chunk: str) -> str:
            return _LATEX_ESCAPE_RE.sub(lambda m: _LATEX_ESCAPE_MAP[m.group(0)], chunk)

        parts = []
        last_end = 0
        for match in _MATH_SPAN_RE.finditer(text):
            parts.append(_escape_plain(text[last_end:match.start()]))
            parts.append(cls._clean_math_spacing(match.group(0)))
            last_end = match.end()
        parts.append(_escape_plain(text[last_end:]))
        return "".join(parts)

    @classmethod
    def _rows_to_latex_table(cls, rows: List[List[str]]) -> str:
        """
        Construye una tabla LaTeX a partir de una lista de filas ya parseadas.

        Usa 'longtable' en vez de 'table[H]' + 'tabularx'. El entorno
        'table' es un *float*: aunque se use la opción [H] del paquete
        'float' para pedirle que se quede "aquí", sigue siendo una unidad
        indivisible que debe caber completa en el espacio libre de la
        página. Si la tabla es más alta que ese espacio, LaTeX no logra
        colocarla y la va empujando hacia adelante; con varias tablas
        seguidas (algo típico en documentos convertidos desde MinerU) el
        cupo de floats sin resolver se agota y todas terminan apareciendo
        juntas al final del documento, en vez de donde corresponden.

        'longtable' no es un float: es una tabla que fluye igual que un
        párrafo normal, exactamente en el punto del documento donde se
        escribe. Si no cabe en lo que resta de la página actual, se corta
        sola justo en el borde de página y continúa automáticamente en la
        siguiente, repitiendo el encabezado (\\endfirsthead/\\endhead) y
        mostrando un aviso de continuación (\\endfoot/\\endlastfoot).
        """
        n_cols = max(len(r) for r in rows)
        norm_rows = [r + [""] * (n_cols - len(r)) for r in rows]

        # Ancho de columna calculado por el propio LaTeX (\dimexpr): reparte
        # \textwidth entre todas las columnas, descontando el padding interno
        # (\tabcolsep, a ambos lados de cada columna) para que la tabla nunca
        # se salga de la caja de texto. Cada columna es 'p{...}' con wrap
        # automático (en vez de 'l', que no ajusta y desborda con celdas largas).
        col_width = f"\\dimexpr(\\textwidth-{2 * n_cols}\\tabcolsep)/{n_cols}\\relax"
        col_spec = (">{\\raggedright\\arraybackslash}p{" + col_width + "}") * n_cols

        header_cells = [cls._escape_latex(cell) for cell in norm_rows[0]]
        header_line = " & ".join(header_cells) + " \\\\"

        lines = []
        lines.append(f"\\begin{{longtable}}[c]{{{col_spec}}}")
        # --- Encabezado de la primera página de la tabla ---
        lines.append("\\toprule")
        lines.append(header_line)
        lines.append("\\midrule")
        lines.append("\\endfirsthead")
        # --- Encabezado repetido en cada página siguiente, si la tabla se corta ---
        lines.append(
            f"\\multicolumn{{{n_cols}}}{{l}}{{\\small\\itshape (continúa de la página anterior)}} \\\\"
        )
        lines.append("\\toprule")
        lines.append(header_line)
        lines.append("\\midrule")
        lines.append("\\endhead")
        # --- Pie que aparece al cortarse la tabla (todas las páginas salvo la última) ---
        lines.append("\\midrule")
        lines.append(
            f"\\multicolumn{{{n_cols}}}{{r}}{{\\small\\itshape (continúa en la página siguiente)}} \\\\"
        )
        lines.append("\\endfoot")
        # --- Pie de la última página de la tabla ---
        lines.append("\\bottomrule")
        lines.append("\\endlastfoot")
        for row in norm_rows[1:]:
            cells = [cls._escape_latex(cell) for cell in row]
            lines.append(" & ".join(cells) + " \\\\")
        lines.append("\\end{longtable}")
        return "\n".join(lines)

    @classmethod
    def _html_table_to_latex(cls, html_content: str) -> str:
        """Convierte una tabla HTML (formato típico de MinerU) a longtable+booktabs."""
        parser = _LatexHTMLTableParser()
        try:
            parser.feed(html_content)
            parser.close()
        except Exception:
            logger.warning("No se pudo parsear una tabla HTML; se exporta como texto plano")
            return cls._escape_latex(re.sub(r"<[^>]+>", " ", html_content))

        rows = [row for row in parser.rows if any(cell.strip() for cell in row)]
        if not rows:
            return cls._escape_latex(re.sub(r"<[^>]+>", " ", html_content))

        return cls._rows_to_latex_table(rows)

    @classmethod
    def _markdown_table_to_latex(cls, content: str) -> str:
        """
        Convierte una tabla en formato markdown (líneas con '|') a longtable+booktabs.

        A veces MinerU/la traducción devuelven la tabla sin saltos de línea
        (todo en un único párrafo, ej: "a | b | | c | d | | e | f |"). En ese
        caso se reconstruyen las filas a partir del patrón '| |', que marca
        el borde entre el final de una fila ("...|") y el inicio de la
        siguiente ("|...").
        """
        normalized = content.strip()
        if "\n" not in normalized:
            normalized = _TABLE_ROW_BOUNDARY_RE.sub("|\n|", normalized)

        rows: List[List[str]] = []
        for line in normalized.splitlines():
            line = line.strip()
            if not line:
                continue
            # Línea separadora markdown típica: |---|---|---| o | :--- | ---: |
            if set(line.replace("|", "").strip()) <= {"-", ":", " "}:
                continue
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if any(cell for cell in cells):
                rows.append(cells)

        # Filtra filas separadoras que usan raya larga "—" en vez de "-"
        # (MinerU/la traducción a veces la sustituye), p. ej. "— | — | —".
        def _is_separator_row(row: List[str]) -> bool:
            return all(set(cell) <= {"-", "—", ":", " "} for cell in row if cell) and any(row)

        rows = [row for row in rows if not _is_separator_row(row)]

        if not rows:
            return cls._escape_latex(content)

        return cls._rows_to_latex_table(rows)

    @classmethod
    def _render_latex_layout_element(cls, element: dict, original_only: bool = False) -> str:
        """Renderiza un elemento estructurado (de layout_elements) a LaTeX válido."""
        element_type = element.get("type", "paragraph")
        content = (
            element.get("content", "")
            if original_only
            else element.get("translated_content") or element.get("content", "")
        )
        if not content.strip():
            return ""

        if element_type == "section_title":
            # Sin numerar: el texto traducido ya trae su propia numeración
            # (p. ej. "1 Introducción"), así que \section normal duplicaría
            # el número. \section* no numera ni la añade al índice.
            return f"\\section*{{{cls._escape_latex(content)}}}"

        if element_type == "table":
            if "<table" in content.lower():
                return cls._html_table_to_latex(content)
            if "|" in content:
                return cls._markdown_table_to_latex(content)
            # Sin '|' ni '<table': no se reconoce estructura, se escapa como texto.
            return cls._escape_latex(content)

        if element_type == "equation":
            inner = content.strip()
            if inner.startswith("$") and inner.endswith("$"):
                return cls._clean_math_spacing(inner)
            return f"\\[{inner}\\]"

        if element_type == "figure":
            return f"% [Figura preservada]\n{cls._escape_latex(content)}"

        # reference / paragraph / cualquier otro tipo
        return cls._escape_latex(content)

    # =========================================================================
    # LaTeX: exportador principal
    # =========================================================================

    def export_to_latex(self, document: Document, output_dir: Optional[Path] = None, original_only: bool = False) -> Path:
        """
        Exporta la traducción preservando formato LaTeX.

        - Si el documento tiene `layout_elements` (extraído por MinerU),
          renderiza cada elemento según su tipo: las tablas HTML se
          convierten a `longtable`/`booktabs` reales (fluyen en el texto y
          se cortan solas entre páginas) y todo el texto se escapa
          correctamente, preservando las fórmulas matemáticas.
        - Si no hay `layout_elements`, cae al camino clásico basado en
          `document.sections`, también con escape de caracteres especiales.
        """
        output_dir = output_dir or self.export_dir
        filename = self._generate_filename(document, suffix="traducido", ext="tex")
        output_path = output_dir / filename

        lines = []
        lines.append("% LaTeX document - Translated by Paper Translator")
        lines.append(f"% Original: {document.filename}")
        lines.append(f"% Date: {datetime.now().strftime('%Y-%m-%d')}")
        lines.append("")
        lines.append("\\documentclass[12pt,a4paper]{article}")
        lines.append("\\usepackage[spanish]{babel}")
        lines.append("\\usepackage[utf8]{inputenc}")
        lines.append("\\usepackage[T1]{fontenc}")
        lines.append("\\usepackage{amsmath,amssymb}")
        lines.append("\\usepackage{graphicx}")
        lines.append("\\usepackage{hyperref}")
        lines.append("\\usepackage{array}")      # habilita columnas p{} con ancho fijo y alineación custom (>{...})
        lines.append("\\usepackage{longtable}")  # tablas que FLUYEN en el texto (no son floats) y cruzan de página solas
        lines.append("\\usepackage{booktabs}")   # \\toprule/\\midrule/\\bottomrule
        lines.append("")
        lines.append("\\title{" + self._escape_latex(Path(document.filename).stem) + "}")
        lines.append("\\author{Traducción automática}")
        lines.append("\\date{\\today}")
        lines.append("")
        lines.append("\\begin{document}")
        lines.append("\\maketitle")
        lines.append("\\clearpage")  # el contenido empieza en una página nueva, tras la portada
        lines.append("")

        # --- Camino "estructurado": usa layout_elements de MinerU si existen ---
        layout_elements = document.metadata.get("layout_elements", [])
        if layout_elements:
            for element in layout_elements:
                rendered = self._render_latex_layout_element(element, original_only=original_only)
                if rendered:
                    lines.append(rendered)
                    lines.append("")
            lines.append("\\end{document}")
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write("\n".join(lines))
            logger.info(f"Exportado LaTeX estructurado: {output_path}")
            return output_path

        # --- Camino "clásico": sin layout_elements, usa document.sections ---
        # Nota: se usan las variantes sin numerar (\section*, \subsection*)
        # porque el texto traducido de section.title ya suele traer su
        # propia numeración (p. ej. "1 Introducción"); si dejáramos que
        # LaTeX numerara también, saldría duplicado ("2. 1 Introducción").
        section_env_map = {
            'abstract': 'abstract',
            'introduction': 'section*',
            'methodology': 'section*',
            'results': 'section*',
            'discussion': 'section*',
            'conclusions': 'section*',
            'references': 'section*',
            'acknowledgments': 'section*',
            'appendix': 'section*',
            'other': 'subsection*',
            'title': 'title',
        }

        for section in document.sections:
            text = (
                section.original_text
                if original_only
                else section.translated_text if section.translated_text else section.original_text
            )
            if not text.strip():
                continue

            stype = section.section_type.value
            env = section_env_map.get(stype, 'section*')

            if stype == 'title':
                continue  # Ya incluido en \title

            escaped_text = self._escape_latex(text)

            if stype == 'abstract':
                lines.append("\\begin{abstract}")
                lines.append(escaped_text)
                lines.append("\\end{abstract}")
            else:
                title = self._escape_latex(section.title or stype.capitalize())
                lines.append(f"\\{env}{{{title}}}")
                lines.append(escaped_text)
            lines.append("")

        lines.append("\\end{document}")

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write("\n".join(lines))

        logger.info(f"Exportado LaTeX: {output_path}")
        return output_path

    def export_original_format(self, document: Document, output_dir: Optional[Path] = None) -> Path:
        """
        Exporta en el formato más apropiado según el tipo de archivo original.
        Ahora incluye PDF como opción para documentos de texto.
        """
        if document.file_type == FileType.PDF:
            return self.export_to_pdf(document, output_dir)  # ✅ Ahora exporta PDF
        elif document.file_type == FileType.DOCX:
            return self.export_to_docx(document, output_dir)
        elif document.file_type == FileType.TXT:
            return self.export_to_txt(document, output_dir)
        elif document.file_type == FileType.LATEX:
            return self.export_to_latex(document, output_dir)
        else:
            return self.export_to_txt(document, output_dir)

    def export_all_formats(
        self,
        document: Document,
        output_dir: Optional[Path] = None,
        formats: List[str] = ["txt", "docx", "pdf"],
        original_only: bool = False,
    ) -> List[Path]:
        """
        Exporta el documento en múltiples formatos a la vez.

        Args:
            document: Documento a exportar
            output_dir: Directorio de salida
            formats: Lista de formatos a exportar: "txt", "docx", "pdf", "latex"

        Returns:
            Lista de rutas de los archivos exportados
        """
        output_dir = output_dir or self.export_dir
        exported = []

        format_map = {
            "txt": self.export_to_txt,
            "docx": self.export_to_docx,
            "pdf": self.export_to_pdf,
            "latex": self.export_to_latex,
        }

        for fmt in formats:
            fmt = fmt.lower().strip()
            if fmt in format_map:
                try:
                    path = format_map[fmt](document, output_dir, original_only=original_only)
                    exported.append(path)
                except Exception as e:
                    logger.error(f"Error exportando a {fmt}: {e}")
            else:
                logger.warning(f"Formato no soportado: {fmt}")

        return exported

    def create_zip_bundle(self, exported_paths: List[Path], zip_name: Optional[str] = None) -> Path:
        """Crea un archivo ZIP con todos los documentos exportados."""
        if not zip_name:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            zip_name = f"traducciones_{timestamp}.zip"

        zip_path = self.export_dir / zip_name

        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
            for path in exported_paths:
                if path.exists():
                    zf.write(str(path), arcname=path.name)

        logger.info(f"ZIP creado: {zip_path} ({len(exported_paths)} archivos)")
        return zip_path