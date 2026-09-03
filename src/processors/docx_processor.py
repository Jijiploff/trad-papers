"""
Procesador para archivos DOCX (Microsoft Word).
Usa python-docx para extraer texto preservando estructura de párrafos.
"""
import io
from typing import Optional

from src.core.models import FileType
from src.processors.base import (
    BaseFileProcessor,
    CorruptedFileError,
    FileProcessingError,
)
from src.utils.logger import get_logger

logger = get_logger(__name__)


class DocxProcessor(BaseFileProcessor):
    """Procesador de archivos DOCX."""

    def can_handle(self, file_type: FileType) -> bool:
        return file_type == FileType.DOCX

    def extract_text(self, document) -> str:
        """Extrae texto de un archivo DOCX preservando párrafos."""
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise FileProcessingError(
                "Se requiere python-docx para procesar archivos DOCX. "
                "Instálalo con: pip install python-docx"
            )

        try:
            docx_bytes = io.BytesIO(document.content)
            doc = DocxDocument(docx_bytes)
        except Exception as e:
            raise CorruptedFileError(f"Archivo DOCX corrupto o ilegible: {str(e)}")

        text_parts = []

        # Extraer párrafos
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                # Preservar encabezados detectando estilos
                style = para.style.name.lower() if para.style else ''
                if 'heading' in style or 'título' in style:
                    text_parts.append(f"\n## {text}\n")
                else:
                    text_parts.append(text)

        # Extraer tablas (convertir a formato texto simple)
        for table in doc.tables:
            table_text = ["\n[Tabla]\n"]
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                table_text.append(" | ".join(cells))
            table_text.append("\n")
            text_parts.append("\n".join(table_text))

        full_text = "\n\n".join(text_parts)

        # Estimar páginas (aproximado: ~500 palabras por página)
        word_count = len(full_text.split())
        self._page_count = max(1, word_count // 500)

        return full_text

    def get_page_count(self, document) -> Optional[int]:
        if hasattr(self, '_page_count'):
            return self._page_count
        return None
